"""Читання Excel-файлів зі списками стипендіатів разом із вбудованими фото."""
from dataclasses import dataclass, field
from typing import Optional

import openpyxl

from .utils import name_key

EMU_PER_PT = 12700
DEFAULT_SPONSOR = "Don Bosco Jugendhilfe Weltweit"  # логотипи в об'єднаній комірці A2:A50


@dataclass
class Row:
    name: str
    group: str = ""
    need: str = ""
    sponsor: str = ""
    photo: Optional[bytes] = None
    source_row: int = 0


@dataclass
class Report:
    conflicts: list = field(default_factory=list)


def _clean(value) -> str:
    return " ".join(str(value or "").split())


def extract_photos(ws, skip_cols=(0,), report: Optional[Report] = None) -> dict:
    """
    Повертає {номер_рядка: bytes}. Фото в Excel «пливуть» по рядках, тому рядок визначається
    за верхнім краєм якоря; якщо верх на 90%+ нижче початку рядка — фото належить наступному рядку.
    """
    default_h = (ws.sheet_format.defaultRowHeight or 15) * EMU_PER_PT
    photos = {}
    for im in ws._images:
        anchor = im.anchor._from
        if anchor.col in skip_cols:
            continue
        row = anchor.row + 1
        height = (ws.row_dimensions[row].height or 0) * EMU_PER_PT or default_h
        if anchor.rowOff >= 0.9 * height:
            row += 1
        data = im._data()
        if row in photos:
            if report is not None:
                report.conflicts.append(row)
            continue
        photos[row] = data
    return photos


def read_tuition(path, report: Optional[Report] = None) -> list[Row]:
    """Файл «Стипендія на навчання»: № | ФОТО | ПІП | професія, група | ..."""
    ws = openpyxl.load_workbook(path).active
    photos = extract_photos(ws, skip_cols=(), report=report)
    rows = []
    for r in range(2, ws.max_row + 1):
        name = _clean(ws.cell(r, 3).value)
        if not name:
            continue
        group = _clean(ws.cell(r, 4).value).split(",")[-1].strip()
        rows.append(Row(name=name, group=group, need="Навчання", photo=photos.get(r), source_row=r))
    return rows


def read_list(path, report: Optional[Report] = None) -> list[Row]:
    """Файл «Список учнів»: Засновник | № | ПІП | ФОТО | Група | Потреба | Портфоліо"""
    ws = openpyxl.load_workbook(path).active
    photos = extract_photos(ws, skip_cols=(0,), report=report)
    rows, sponsor = [], DEFAULT_SPONSOR
    for r in range(2, ws.max_row + 1):
        founder = ws.cell(r, 1).value
        if founder and str(founder).strip():
            sponsor = _clean(str(founder).split("\n")[0])
        name = _clean(ws.cell(r, 3).value)
        if not name:
            continue
        rows.append(Row(
            name=name,
            group=_clean(ws.cell(r, 5).value),
            need=_clean(ws.cell(r, 6).value),
            sponsor=sponsor,
            photo=photos.get(r),
            source_row=r,
        ))
    return rows


def kinds_from_need(need: str) -> set:
    n = need.casefold()
    kinds = set()
    if "колегіум" in n:
        kinds.add("collegium")
    if "харчуван" in n:
        kinds.add("cafe")
    if "навчан" in n:
        kinds.add("tuition")
    return kinds


def merge(tuition_rows: list[Row], list_rows: list[Row]) -> dict:
    """Об'єднує обидва файли за ПІП: {ключ: {...}}."""
    people = {}
    for row in list_rows:
        p = people.setdefault(name_key(row.name), {
            "name": row.name, "group": row.group, "kinds": set(), "sponsor_kinds": set(),
            "sponsor": row.sponsor, "photo": row.photo,
        })
        kinds = kinds_from_need(row.need)
        p["kinds"] |= kinds
        p["sponsor_kinds"] |= kinds
        p["photo"] = p["photo"] or row.photo
    for row in tuition_rows:
        p = people.setdefault(name_key(row.name), {
            "name": row.name, "group": row.group, "kinds": set(), "sponsor_kinds": set(),
            "sponsor": "", "photo": row.photo,
        })
        p["kinds"].add("tuition")
        p["group"] = p["group"] or row.group
        p["photo"] = p["photo"] or row.photo
    return people
